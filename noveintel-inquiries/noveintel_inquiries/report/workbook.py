"""
The companion workbook — the run's cases, declared and identity-free (R3, R10).

Column order is declared, not a per-flow constant: the universal columns the tenant
actually sent, in U-id order, then the derived columns D1–D8 and the audit fields.
No sensitive column exists here to write — identity was excluded at ingest — and the
free text is the redacted text the classifier saw.

Sheets: a summary; every case; one sheet per actual contact type (D1); the
reclassification sheet (D4); the deep-reviewed cases; and the Domain 02 handoff — cases
whose true type is a complaint, incident or compliment. Every case appears in exactly
one per-type sheet, so the per-type sheets sum to the cover total (R2).
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from ..models.columns import COLUMN_BY_NAME
from ..engine.parse import Case
from ..engine.wording import CONTACT_TYPES, PROVENANCE, TIMELINESS_LABELS

__all__ = ["write_workbook", "WORKBOOK_DERIVED_COLUMNS"]

_UNIVERSAL_ATTR = {
    "case_id": "case_id", "description": "description", "resolution_text": "resolution_text",
    "case_type": "case_type_raw", "crm_service": "crm_service", "crm_service_group": "crm_service_group",
    "severity": "severity", "channel": "channel", "opened_datetime": "opened", "closed_datetime": "closed",
    "sla_status": "sla", "status": "status", "department": "department", "owner": "owner",
    "resolved_by": "resolved_by", "sub_classification_hint": "sub_hint", "escalation_datetime": "escalation",
    "entity": "entity", "accepted_datetime": "accepted", "classified_datetime": "classified_at",
    "first_response_datetime": "first_response", "resolved_datetime": "resolved_at",
    "customer_verdict": "customer_verdict", "reopen_flag": "reopen", "fcr_flag": "fcr",
    "closure_reason": "closure_reason", "satisfaction_score": "satisfaction", "crm_title": "crm_title",
    "updated_datetime": "updated",
}

WORKBOOK_DERIVED_COLUMNS: list[tuple[str, str]] = [
    ("D1", "التصنيف الفعلي"),
    ("D1b", "النوع الحقيقي"),
    ("D2a", "الخدمة الرئيسية"),
    ("D2", "التصنيف الفرعي"),
    ("D3", "سبب التصنيف"),
    ("D4", "إعادة التصنيف"),
    ("D5", "ثقة التصنيف"),
    ("D6", "مصدر التصنيف"),
    ("D7", "السبب الجذري"),
    ("D8", "الحزمة الموضوعية"),
    ("REV", "خضعت للمراجعة المعمّقة"),
    ("LOW", "ثقة منخفضة بعد المراجعة"),
    ("DUP", "منسوخة من حالة مطابقة"),
    ("FAQ", "السؤال الشائع المرتبط"),
    ("FIX", "التصحيحات الآلية"),
]


def _value(case: Case, column: str) -> Any:
    value = getattr(case, _UNIVERSAL_ATTR[column])
    if column == "sla_status":
        return TIMELINESS_LABELS.get(value, {}).get("ar", value)
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, bool):
        return "نعم" if value else "لا"
    return value


def _derived(case: Case, umbrella_names: dict[str, str]) -> list[Any]:
    misclassified = case.misclassified
    return [
        case.contact_type or "",
        case.true_type or "",
        case.main_service or "",
        case.sub_classification or "",
        case.reason or case.evidence or "",
        "" if misclassified is None else ("نعم" if misclassified else "لا"),
        round(case.confidence, 3) if case.confidence is not None else "",
        PROVENANCE.get(case.provenance or "", case.provenance or ""),
        case.root_cause or "",
        umbrella_names.get(case.bundle or "", case.bundle or ""),
        "نعم" if case.deep_reviewed else "لا",
        "نعم" if case.low_confidence else "لا",
        case.duplicate_of or "",
        case.faq_id or "",
        "؛ ".join(case.corrections),
    ]


def write_workbook(path: Path, cases: list[Case], columns: list[str], metrics: dict[str, Any],
                   umbrella_names: dict[str, str], primary: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor=primary.lstrip("#").upper())
    header_font = Font(bold=True, color="FFFFFF")
    universal = [c for c in columns if c in _UNIVERSAL_ATTR]
    headers = [f"{COLUMN_BY_NAME[c].uid} {COLUMN_BY_NAME[c].label_ar}" for c in universal] + [
        f"{uid} {label}" for uid, label in WORKBOOK_DERIVED_COLUMNS
    ]

    wb = Workbook()
    summary = wb.active
    summary.title = "ملخص"
    summary.sheet_view.rightToLeft = True

    def add_sheet(title: str, rows: list[Case]) -> None:
        ws = wb.create_sheet(title[:31])
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for cell in ws[1]:
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for case in rows:
            ws.append([_value(case, c) for c in universal] + _derived(case, umbrella_names))
        ws.freeze_panes = "A2"
        for index, _ in enumerate(headers, 1):
            ws.column_dimensions[ws.cell(1, index).column_letter].width = 22

    kpi = metrics["KPI"]
    rows = [
        ("المؤشر", "القيمة"),
        ("إجمالي الحالات ضمن النطاق", kpi["total_cases"]),
        ("الحالات المصنفة", kpi["classified"]),
        ("الحالات المغلقة", kpi.get("closed_cases") if kpi.get("closed_cases") is not None else "لم يُزوَّد"),
        ("نسبة إعادة التصنيف", kpi.get("reclassification_rate")),
        ("الالتزام بزمن الاستجابة (ضمن المهلة) %", kpi.get("within_deadline_pct")),
        ("نسبة القنوات الرقمية %", kpi.get("digital_share_pct")),
        ("حالات خضعت للمراجعة المعمّقة", kpi.get("deep_reviewed")),
        ("حالات بثقة منخفضة بعد المراجعة", kpi.get("low_confidence_after_review")),
        ("", ""),
        ("التصنيف الفعلي", "العدد"),
    ]
    rows += list(zip(metrics["ACTUAL_TYPE"]["labels"], metrics["ACTUAL_TYPE"]["values"]))
    rows += [("", ""), ("مصدر التصنيف", "العدد")]
    rows += list(metrics["CLASSIFICATION_QUALITY"]["provenance"].items())
    if metrics.get("CHANNEL", {}).get("labels"):
        rows += [("", ""), ("القناة", "العدد")]
        rows += list(zip(metrics["CHANNEL"]["labels"], metrics["CHANNEL"]["values"]))
    for row in rows:
        summary.append(list(row))
    for cell in summary[1]:
        cell.fill, cell.font = header_fill, header_font
    summary.column_dimensions["A"].width = 44
    summary.column_dimensions["B"].width = 18

    add_sheet("كل الحالات", cases)
    sheet_names = {"استفسار": "استفسارات", "طلب خدمة": "طلبات خدمة", "شكوى": "شكاوى", "شكر وثناء": "شكر وثناء"}
    for contact in CONTACT_TYPES:
        add_sheet(sheet_names[contact], [c for c in cases if c.contact_type == contact])
    unlabelled = [c for c in cases if c.contact_type not in CONTACT_TYPES]
    if unlabelled:
        add_sheet("غير مصنف", unlabelled)
    add_sheet("إعادة التصنيف", [c for c in cases if c.misclassified])
    add_sheet("المراجعة المعمّقة", [c for c in cases if c.deep_reviewed])
    add_sheet("خارج النطاق - للشكاوى", [c for c in cases if c.true_type in ("complaint", "incident", "compliment")])
    wb.save(path)
