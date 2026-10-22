"""Shared fixtures. No network: the store is an embedded Postgres, the embedder hashes."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import pytest
from openpyxl import Workbook


@pytest.fixture(scope="session")
def pg_dsn(tmp_path_factory):
    pgserver = pytest.importorskip("pgserver", reason="the embedded Postgres (pgserver) is not installed")
    server = pgserver.get_server(tmp_path_factory.mktemp("pg"), cleanup_mode="delete")
    yield server.get_uri()


@pytest.fixture
def store(pg_dsn):
    from noveintel_ingest.store import PgVectorStore

    with PgVectorStore(pg_dsn) as s:
        yield s
        for c in s.collections():
            s.drop_collection(c["name"])


def make_faq_workbook(path: Path, rows: list[tuple] | None = None) -> Path:
    """A FAQ workbook shaped like the ADPF one: a visible list and a hidden draft."""
    wb = Workbook()
    ws = wb.active
    ws.title = "FAQs"
    ws.append(["CA Reviewed Arabic Question", "الصيغة المقترحة ", "الإدارة المعنية", "الخدمة"])
    for row in rows or [
        ("كيف يمكن للمؤمن عليه متابعة طلبه؟", "يمكن متابعة حالة الطلب عبر منصة تم.", "معاشات المدنيين", "صرف المعاش التقاعدي"),
        ("ما هو نظام الخصم المباشر؟", "نظام يهدف إلى تسهيل سداد الاشتراكات الشهرية.", "التحصيل والمتابعة", "تسجيل جهة عمل"),
        ("هل يجوز إعفاء جهات العمل من المبالغ الإضافية؟", "لا يجوز، ويحق لجهة العمل طلب التقسيط.", "التحصيل والمتابعة", "تقسيط"),
    ]:
        ws.append(list(row))
    draft = wb.create_sheet("مسودة")
    draft.sheet_state = "hidden"
    draft.append(["#", "CA Reviewed English Question", "CA Reviewed Arabic Question", "CA Reviewed Arabic Answer"])
    draft.append([6, "Can I re-register?", "هل يمكنني إعادة التسجيل؟", "نعم بعد سنة."])
    wb.save(path)
    return path


@pytest.fixture
def faq_xlsx(tmp_path) -> Path:
    return make_faq_workbook(tmp_path / "الاسئلة الشائعة.xlsx")


@pytest.fixture
def messy_xlsx(tmp_path) -> Path:
    """A title above the table, a blank row, merged cells, a blank and a repeated header."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws["A1"] = "Quarterly orders report"
    ws.merge_cells("A1:E1")
    # row 2 blank
    ws.append([])
    ws.append(["Order ID", "Date", "Region", None, "Region"])
    ws.append([42, dt.datetime(2024, 3, 1), "North", 18500.0, "N1"])
    ws.append([43, dt.datetime(2024, 3, 2, 14, 30), "South", 99.5, "S1"])
    ws.append([None, None, None, None, None])
    ws.append(["Order ID", "Date", "Region", None, "Region"])  # header repeated by a paged export
    ws.append([44, dt.date(2024, 3, 3), "  East  ", 7, "E1"])
    ws.merge_cells("C8:C9")
    ws.append([45, dt.date(2024, 3, 4), None, 8, "E2"])
    other = wb.create_sheet("Hidden notes")
    other.sheet_state = "hidden"
    other.append(["Note"])
    other.append(["secret-ish"])
    wb.save(tmp_path / "orders.xlsx")
    return tmp_path / "orders.xlsx"
