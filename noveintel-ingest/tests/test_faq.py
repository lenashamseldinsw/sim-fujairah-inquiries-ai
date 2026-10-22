import pytest
from openpyxl import Workbook

from noveintel_ingest.faq import faq_catalogue

# The inquiries service's FaqEntry fields (api/models/reference.py). Not imported:
# this package depends on neither service.
FAQ_ENTRY_FIELDS = {"question_ar", "answer_ar", "question_en", "answer_en", "main_service", "keywords"}


def test_visible_sheet_preferred_and_columns_found_by_header(faq_xlsx):
    entries, note = faq_catalogue(faq_xlsx)
    assert note["sheet"] == "FAQs" and not note["sheet_hidden"]
    assert note["columns"] == {"question_ar": "CA Reviewed Arabic Question", "answer_ar": "الصيغة المقترحة"}
    assert len(entries) == 3
    assert entries[0] == {
        "question_ar": "كيف يمكن للمؤمن عليه متابعة طلبه؟",
        "answer_ar": "يمكن متابعة حالة الطلب عبر منصة تم.",
        "keywords": [],
    }
    assert all(set(e) <= FAQ_ENTRY_FIELDS for e in entries)


def test_service_column_and_a_named_sheet(faq_xlsx):
    entries, _ = faq_catalogue(faq_xlsx, main_service="الخدمة")
    assert entries[0]["main_service"] == "صرف المعاش التقاعدي"
    draft, note = faq_catalogue(faq_xlsx, sheet="مسودة")
    assert note["columns"]["question_en"] == "CA Reviewed English Question"
    assert draft == [{"question_en": "Can I re-register?", "question_ar": "هل يمكنني إعادة التسجيل؟",
                      "answer_ar": "نعم بعد سنة.", "keywords": []}]


def test_language_read_from_values_when_header_does_not_say(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Question", "Answer", "السؤال", "الإجابة"])
    ws.append(["How do I track it?", "On TAMM.", "كيف أتابعه؟", "عبر تم."])
    wb.save(tmp_path / "bilingual.xlsx")
    entries, note = faq_catalogue(tmp_path / "bilingual.xlsx")
    assert note["columns"] == {"question_en": "Question", "answer_en": "Answer", "question_ar": "السؤال", "answer_ar": "الإجابة"}
    assert entries[0]["answer_en"] == "On TAMM."


def test_ambiguous_or_missing_columns_refused(tmp_path):
    wb = Workbook()
    wb.active.append(["السؤال", "الإجابة"])
    wb.active.append(["س١", "ج١"])
    second = wb.create_sheet("Other")
    second.append(["السؤال", "الإجابة"])
    second.append(["س٢", "ج٢"])
    wb.save(tmp_path / "two.xlsx")
    with pytest.raises(ValueError, match="Several sheets"):
        faq_catalogue(tmp_path / "two.xlsx")
    with pytest.raises(ValueError, match="no column"):
        faq_catalogue(tmp_path / "two.xlsx", sheet="Other", answer_ar="Answer")
